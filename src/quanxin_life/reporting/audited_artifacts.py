"""Deterministic artifacts rendered only from ledger-registered audited reports."""

from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass
from enum import StrEnum
from hashlib import sha256 as hash_sha256
from pathlib import Path
from typing import Any, Protocol
from uuid import UUID

from pydantic import field_validator

from quanxin_life.core import ToolResult
from quanxin_life.core.schemas import ContractModel, Sha256
from quanxin_life.reporting.audited_markdown import REPORTING_VERSION
from quanxin_life.reporting.contracts import (
    AUDITED_REPORT_TOOL_NAME,
    AUDITED_REPORT_TOOL_VERSION,
    RECOMMENDATION_REPORT_RENDERER_VERSION,
)

_SUPPORTED_REPORT_RENDERERS = frozenset(
    {REPORTING_VERSION, RECOMMENDATION_REPORT_RENDERER_VERSION}
)


class ReportExportDependencyUnavailable(RuntimeError):
    """Raised only when an explicitly requested optional renderer is absent."""


class ReportArtifactFormat(StrEnum):
    JSON = "json"
    MARKDOWN = "markdown"
    PDF = "pdf"
    DOCX = "docx"


class ReviewedPdfFont(ContractModel):
    """Operator-owned font whose exact bytes are reviewed before PDF rendering."""

    path: Path
    sha256: Sha256

    @field_validator("path")
    @classmethod
    def path_is_absolute(cls, value: Path) -> Path:
        if not value.is_absolute():
            raise ValueError("PDF font path must be absolute")
        return value

    def verified_path(self) -> Path:
        if not self.path.is_file():
            raise ValueError("reviewed PDF font file was not found")
        actual = hash_sha256(self.path.read_bytes()).hexdigest()
        if actual != self.sha256:
            raise ValueError("reviewed PDF font SHA-256 does not match")
        return self.path


REPORTLAB_VERA_FONT_SHA256 = (
    "c4c45690b345435b2cba52ecabe275f05e49b389b39fe68ad03afbb551288d3d"
)


def reviewed_reportlab_vera_font() -> ReviewedPdfFont:
    """Locate ReportLab's pinned, SHA-reviewed portable TrueType font."""

    try:
        import reportlab  # type: ignore[import-untyped]
    except ModuleNotFoundError as exc:  # pragma: no cover - optional dependency path
        raise ReportExportDependencyUnavailable(
            "PDF export requires installing the reporting dependency group"
        ) from exc
    path = Path(reportlab.__file__).resolve().parent / "fonts" / "Vera.ttf"
    font = ReviewedPdfFont(path=path, sha256=REPORTLAB_VERA_FONT_SHA256)
    font.verified_path()
    return font


@dataclass(frozen=True, slots=True)
class AuditedReportArtifact:
    """One detached artifact with transport metadata and a byte hash."""

    source_result_id: str
    format: ReportArtifactFormat
    filename: str
    media_type: str
    payload: bytes
    sha256: str


class RegisteredResultResolver(Protocol):
    def resolve_registered_result(self, result_id: str) -> ToolResult: ...


class AuditedReportArtifactExporter:
    """Resolve one audited result from the ledger and render approved formats."""

    def __init__(
        self,
        ledger: RegisteredResultResolver,
        *,
        pdf_font: ReviewedPdfFont | None = None,
    ) -> None:
        self._ledger = ledger
        self._pdf_font = pdf_font

    def export(
        self,
        result_id: str,
        format: ReportArtifactFormat,
    ) -> AuditedReportArtifact:
        result = self._resolve_report(result_id)
        report_id, markdown = _report_content(result)
        if format is ReportArtifactFormat.JSON:
            payload = _render_json(result)
            extension = "json"
            media_type = "application/json"
        elif format is ReportArtifactFormat.MARKDOWN:
            payload = markdown.encode("utf-8")
            extension = "md"
            media_type = "text/markdown; charset=utf-8"
        elif format is ReportArtifactFormat.DOCX:
            payload = _render_docx(markdown, result=result)
            extension = "docx"
            media_type = (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        elif format is ReportArtifactFormat.PDF:
            if self._pdf_font is None:
                raise ValueError("PDF export requires a reviewed font")
            payload = _render_pdf(
                markdown,
                font=self._pdf_font,
                source_result_id=result.result_id,
            )
            extension = "pdf"
            media_type = "application/pdf"
        else:  # pragma: no cover - exhaustive StrEnum branch
            raise ValueError("unsupported audited report artifact format")
        return AuditedReportArtifact(
            source_result_id=result.result_id,
            format=format,
            filename=f"audited-report-{report_id}.{extension}",
            media_type=media_type,
            payload=payload,
            sha256=hash_sha256(payload).hexdigest(),
        )

    def _resolve_report(self, result_id: str) -> ToolResult:
        result = self._ledger.resolve_registered_result(result_id)
        if (
            result.tool_name != AUDITED_REPORT_TOOL_NAME
            or result.tool_version != AUDITED_REPORT_TOOL_VERSION
            or result.model_version not in _SUPPORTED_REPORT_RENDERERS
        ):
            raise ValueError("artifact source must be a supported audited report")
        return result


def _report_content(result: ToolResult) -> tuple[str, str]:
    report_id = result.values.get("report_id")
    markdown = result.values.get("markdown")
    rendering_version = result.values.get("rendering_version")
    if not isinstance(report_id, str) or not report_id.strip():
        raise ValueError("audited report result contains an invalid report_id")
    try:
        UUID(report_id)
    except (AttributeError, TypeError, ValueError) as exc:
        raise ValueError("audited report report_id must be a UUID") from exc
    if not isinstance(markdown, str) or not markdown.strip():
        raise ValueError("audited report result contains no Markdown")
    if (
        rendering_version not in _SUPPORTED_REPORT_RENDERERS
        or rendering_version != result.model_version
    ):
        raise ValueError("audited report rendering version is unsupported")
    return report_id, markdown


def _render_json(result: ToolResult) -> bytes:
    return (
        json.dumps(
            result.model_dump(mode="json"),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        )
        + "\n"
    ).encode("utf-8")


def _markdown_blocks(markdown: str) -> tuple[tuple[str, str], ...]:
    blocks: list[tuple[str, str]] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("heading", line[3:].strip()))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].replace("`", "")))
        else:
            blocks.append(("body", line.replace("`", "")))
    if not blocks or blocks[0][0] != "title":
        raise ValueError("audited Markdown must begin with one level-one title")
    return tuple(blocks)


def _render_docx(markdown: str, *, result: ToolResult) -> bytes:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION_START
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ModuleNotFoundError as exc:
        raise ReportExportDependencyUnavailable(
            "DOCX export requires installing the 'quanxin-life[reporting]' extra"
        ) from exc

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    def set_east_asia_font(element: Any) -> None:
        element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Microsoft YaHei"
        )

    set_east_asia_font(normal._element)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    heading = document.styles["Heading 1"]
    heading.font.name = "Calibri"
    heading.font.size = Pt(16)
    heading.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    set_east_asia_font(heading._element)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "Hiro | Audited Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        set_east_asia_font(run._element)

    for kind, text in _markdown_blocks(markdown):
        if kind == "title":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(16)
            run = paragraph.add_run(text)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(23)
            set_east_asia_font(run._element)
        elif kind == "heading":
            document.add_paragraph(text, style="Heading 1")
        elif kind == "bullet":
            paragraph = document.add_paragraph(text, style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing = 1.167
        else:
            document.add_paragraph(text)

    properties = document.core_properties
    properties.title = _markdown_blocks(markdown)[0][1]
    properties.subject = "Ledger-bound audited battery report"
    properties.author = "Hiro"
    properties.identifier = result.result_id
    properties.created = result.created_at.replace(tzinfo=None)
    properties.modified = result.created_at.replace(tzinfo=None)
    raw = io.BytesIO()
    document.save(raw)
    return _normalize_zip(raw.getvalue())


def _normalize_zip(payload: bytes) -> bytes:
    source = io.BytesIO(payload)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(source, "r") as archive,
        zipfile.ZipFile(
            output,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as normalized,
    ):
        for name in sorted(archive.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            normalized.writestr(info, archive.read(name))
    return output.getvalue()


def _render_pdf(
    markdown: str,
    *,
    font: ReviewedPdfFont,
    source_result_id: str,
) -> bytes:
    try:
        from reportlab.lib import colors  # type: ignore[import-untyped]
        from reportlab.lib.enums import TA_RIGHT  # type: ignore[import-untyped]
        from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
        from reportlab.lib.styles import ParagraphStyle  # type: ignore[import-untyped]
        from reportlab.lib.units import inch  # type: ignore[import-untyped]
        from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
        from reportlab.pdfbase.ttfonts import TTFont  # type: ignore[import-untyped]
        from reportlab.pdfgen import canvas  # type: ignore[import-untyped]
        from reportlab.platypus import (  # type: ignore[import-untyped]
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ModuleNotFoundError as exc:
        raise ReportExportDependencyUnavailable(
            "PDF export requires installing the 'quanxin-life[reporting]' extra"
        ) from exc

    from xml.sax.saxutils import escape

    font_path = font.verified_path()
    font_name = f"QuanxinReviewed-{font.sha256[:12]}"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=_markdown_blocks(markdown)[0][1],
        author="Hiro",
        subject="Ledger-bound audited battery report",
    )
    styles = {
        "title": ParagraphStyle(
            "AuditedTitle",
            fontName=font_name,
            fontSize=23,
            leading=28,
            textColor=colors.black,
            spaceAfter=16,
        ),
        "heading": ParagraphStyle(
            "AuditedHeading",
            fontName=font_name,
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=16,
            spaceAfter=8,
        ),
        "body": ParagraphStyle(
            "AuditedBody",
            fontName=font_name,
            fontSize=11,
            leading=14,
            textColor=colors.black,
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "AuditedBullet",
            fontName=font_name,
            fontSize=11,
            leading=14,
            leftIndent=36,
            firstLineIndent=-18,
            bulletIndent=18,
            spaceAfter=8,
        ),
        "footer": ParagraphStyle(
            "AuditedFooter",
            fontName=font_name,
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#666666"),
            alignment=TA_RIGHT,
        ),
    }
    story: list[Any] = []
    for kind, text in _markdown_blocks(markdown):
        safe_text = escape(text)
        if kind == "bullet":
            story.append(Paragraph(safe_text, styles[kind], bulletText="•"))
        else:
            story.append(Paragraph(safe_text, styles[kind]))
        if kind == "title":
            story.append(Spacer(1, 4))

    def canvas_maker(*args: Any, **kwargs: Any) -> Any:
        kwargs["invariant"] = 1
        return canvas.Canvas(*args, **kwargs)

    def add_footer(pdf_canvas: Any, _: Any) -> None:
        pdf_canvas.saveState()
        footer = Paragraph(
            f"Audited ToolResult: {escape(source_result_id)}",
            styles["footer"],
        )
        footer.wrapOn(pdf_canvas, 6.5 * inch, 0.3 * inch)
        footer.drawOn(pdf_canvas, inch, 0.45 * inch)
        pdf_canvas.restoreState()

    document.build(
        story,
        onFirstPage=add_footer,
        onLaterPages=add_footer,
        canvasmaker=canvas_maker,
    )
    return buffer.getvalue()


__all__ = [
    "REPORTLAB_VERA_FONT_SHA256",
    "AuditedReportArtifact",
    "AuditedReportArtifactExporter",
    "ReportArtifactFormat",
    "ReportExportDependencyUnavailable",
    "ReviewedPdfFont",
    "reviewed_reportlab_vera_font",
]
