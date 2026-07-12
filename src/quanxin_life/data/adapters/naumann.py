"""Explicit-layout adapter for provenance-verified Naumann workbooks.

The official workbooks are not present in this repository.  Callers must
therefore supply a reviewed layout instead of relying on guessed sheet names,
headers, units, or capacity defaults.
"""

import math
from collections.abc import Iterable
from pathlib import Path
from typing import Literal, TypeAlias, cast

from pydantic import BaseModel, ConfigDict, Field, model_validator

from quanxin_life.core import CellMetadata
from quanxin_life.data.manifest import RawFileManifest, verify_raw_file_stream
from quanxin_life.data.schemas import CycleRecord, DataQualitySeverity
from quanxin_life.data.source_catalog import IngestionMode, SourceCatalogEntry
from quanxin_life.data.validation import validate_cycle_records

NaumannDatasetId = Literal["NAUMANN_CYCLE", "NAUMANN_CALENDAR"]
NaumannCell: TypeAlias = tuple[CellMetadata, tuple[CycleRecord, ...]]

_REQUIRED_UNITS = {
    "nominal_capacity_ah": "Ah",
    "time_s": "s",
    "voltage_v": "V",
    "current_a": "A",
}
_OPTIONAL_UNITS = {
    "temperature_c": "degC",
    "charge_capacity_ah": "Ah",
    "discharge_capacity_ah": "Ah",
    "internal_resistance_ohm": "ohm",
}
_ADAPTER_VERSION = "naumann-explicit-layout-v1.0.1"


class NaumannWorkbookUnits(BaseModel):
    """Canonical units declared by a reviewed Naumann workbook layout."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    nominal_capacity_ah: Literal["Ah"]
    time_s: Literal["s"]
    voltage_v: Literal["V"]
    current_a: Literal["A"]
    temperature_c: Literal["degC"] | None = None
    charge_capacity_ah: Literal["Ah"] | None = None
    discharge_capacity_ah: Literal["Ah"] | None = None
    internal_resistance_ohm: Literal["ohm"] | None = None


class NaumannWorkbookLayout(BaseModel):
    """A reviewed workbook schema; no header aliases or inferred units exist."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    layout_version: str = Field(min_length=1)
    sheet_name: str = Field(min_length=1)
    header_row: int = Field(ge=1)
    raw_cell_id_column: str = Field(min_length=1)
    nominal_capacity_ah_column: str = Field(min_length=1)
    cycle_index_column: str = Field(min_length=1)
    sample_index_column: str = Field(min_length=1)
    time_s_column: str = Field(min_length=1)
    voltage_v_column: str = Field(min_length=1)
    current_a_column: str = Field(min_length=1)
    temperature_c_column: str | None = None
    charge_capacity_ah_column: str | None = None
    discharge_capacity_ah_column: str | None = None
    internal_resistance_ohm_column: str | None = None
    protocol_id_column: str | None = None
    protocol_description_column: str | None = None
    units: NaumannWorkbookUnits

    @model_validator(mode="after")
    def columns_and_units_are_explicit(self) -> "NaumannWorkbookLayout":
        columns = [
            self.raw_cell_id_column,
            self.nominal_capacity_ah_column,
            self.cycle_index_column,
            self.sample_index_column,
            self.time_s_column,
            self.voltage_v_column,
            self.current_a_column,
            self.temperature_c_column,
            self.charge_capacity_ah_column,
            self.discharge_capacity_ah_column,
            self.internal_resistance_ohm_column,
            self.protocol_id_column,
            self.protocol_description_column,
        ]
        present = [column for column in columns if column is not None]
        if len(present) != len(set(present)):
            raise ValueError("layout column names must be unique")
        required_units = (
            ("nominal_capacity_ah", self.units.nominal_capacity_ah),
            ("time_s", self.units.time_s),
            ("voltage_v", self.units.voltage_v),
            ("current_a", self.units.current_a),
        )
        for field, declared_unit in required_units:
            required_expected_unit = _REQUIRED_UNITS[field]
            if declared_unit != required_expected_unit:
                raise ValueError(f"{field} must declare canonical unit {required_expected_unit}")
        optional_columns: tuple[tuple[str, str | None, str | None], ...] = (
            ("temperature_c", self.temperature_c_column, self.units.temperature_c),
            (
                "charge_capacity_ah",
                self.charge_capacity_ah_column,
                self.units.charge_capacity_ah,
            ),
            (
                "discharge_capacity_ah",
                self.discharge_capacity_ah_column,
                self.units.discharge_capacity_ah,
            ),
            (
                "internal_resistance_ohm",
                self.internal_resistance_ohm_column,
                self.units.internal_resistance_ohm,
            ),
        )
        for optional_field, column, optional_declared_unit in optional_columns:
            optional_expected_unit = _OPTIONAL_UNITS[optional_field]
            if column is not None and optional_declared_unit != optional_expected_unit:
                raise ValueError(
                    f"{optional_field} must declare canonical unit {optional_expected_unit}"
                )
        return self


def _require_text(value: object, *, column: str, row_number: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"row {row_number} requires non-empty text in column {column!r}")
    return value.strip()


def _require_finite_float(value: object, *, column: str, row_number: int) -> float:
    if isinstance(value, bool):
        raise ValueError(f"row {row_number} has invalid numeric value in column {column!r}")
    try:
        number = float(cast(str | float | int, value))
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"row {row_number} has invalid numeric value in column {column!r}"
        ) from exc
    if not math.isfinite(number):
        raise ValueError(f"row {row_number} has non-finite value in column {column!r}")
    return number


def _require_index(value: object, *, column: str, row_number: int) -> int:
    number = _require_finite_float(value, column=column, row_number=row_number)
    if number < 0 or not number.is_integer():
        raise ValueError(f"row {row_number} requires a non-negative integer in column {column!r}")
    return int(number)


def _optional_float(value: object, *, column: str, row_number: int) -> float | None:
    if value is None or value == "":
        return None
    return _require_finite_float(value, column=column, row_number=row_number)


def _optional_text(value: object, *, column: str, row_number: int) -> str | None:
    if value is None or value == "":
        return None
    return _require_text(value, column=column, row_number=row_number)


def _column_values(layout: NaumannWorkbookLayout) -> tuple[str, ...]:
    columns = (
        layout.raw_cell_id_column,
        layout.nominal_capacity_ah_column,
        layout.cycle_index_column,
        layout.sample_index_column,
        layout.time_s_column,
        layout.voltage_v_column,
        layout.current_a_column,
        layout.temperature_c_column,
        layout.charge_capacity_ah_column,
        layout.discharge_capacity_ah_column,
        layout.internal_resistance_ohm_column,
        layout.protocol_id_column,
        layout.protocol_description_column,
    )
    return tuple(column for column in columns if column is not None)


def _validate_source(
    path: Path, manifest: RawFileManifest, source: SourceCatalogEntry
) -> NaumannDatasetId:
    allowed = {"NAUMANN_CYCLE", "NAUMANN_CALENDAR"}
    if manifest.dataset_id not in allowed or source.dataset_id != manifest.dataset_id:
        raise ValueError("Naumann manifest and source dataset_id must match a supported dataset")
    if source.ingestion_mode != IngestionMode.TABULAR:
        raise ValueError("Naumann source must use tabular ingestion")
    if path.suffix.lower() != ".xlsx" or ".xlsx" not in {
        suffix.lower() for suffix in source.expected_suffixes
    }:
        raise ValueError("Naumann workbook must use an approved .xlsx suffix")
    if manifest.source_uri != source.source_uri:
        raise ValueError("Naumann manifest source URI does not match the source catalog")
    if manifest.license_name != source.license_status:
        raise ValueError("Naumann manifest license does not match the source catalog")
    return manifest.dataset_id  # type: ignore[return-value]


def _header_mapping(
    header_values: Iterable[object], layout: NaumannWorkbookLayout
) -> dict[str, int]:
    headers: list[str] = []
    for value in header_values:
        if not isinstance(value, str) or not value.strip():
            headers.append("")
        else:
            headers.append(value.strip())
    non_empty = [header for header in headers if header]
    if len(non_empty) != len(set(non_empty)):
        raise ValueError("Naumann workbook header contains duplicate column names")
    mapping = {header: index for index, header in enumerate(headers) if header}
    missing = [column for column in _column_values(layout) if column not in mapping]
    if missing:
        raise ValueError("Naumann workbook is missing required columns: " + ", ".join(missing))
    return mapping


def load_naumann_workbook(
    path: Path,
    manifest: RawFileManifest,
    source: SourceCatalogEntry,
    *,
    layout: NaumannWorkbookLayout,
) -> tuple[NaumannCell, ...]:
    """Load a workbook only through an explicit, source-reviewed layout."""
    path = Path(path)
    layout = NaumannWorkbookLayout.model_validate(layout.model_dump(warnings="none"))
    dataset_id = _validate_source(path, manifest, source)
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - only without the data extra
        raise RuntimeError("Naumann loading requires the 'data' optional dependencies") from exc

    with path.open("rb") as workbook_handle:
        source_sha256 = verify_raw_file_stream(workbook_handle, path, manifest)
        workbook = openpyxl.load_workbook(workbook_handle, read_only=True, data_only=True)
        try:
            if layout.sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Naumann workbook is missing reviewed sheet {layout.sheet_name!r}"
                )
            worksheet = workbook[layout.sheet_name]
            header_values = next(
                worksheet.iter_rows(
                    min_row=layout.header_row,
                    max_row=layout.header_row,
                    values_only=True,
                ),
                None,
            )
            if header_values is None:
                raise ValueError("Naumann workbook is missing the configured header row")
            columns = _header_mapping(header_values, layout)

            records_by_raw_id: dict[str, list[CycleRecord]] = {}
            capacity_by_raw_id: dict[str, float] = {}
            protocol_by_raw_id: dict[str, tuple[str | None, str | None]] = {}
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=layout.header_row + 1, values_only=True),
                start=layout.header_row + 1,
            ):
                if not any(value is not None and value != "" for value in row):
                    continue
                values = {
                    name: row[index] if index < len(row) else None
                    for name, index in columns.items()
                }
                raw_cell_id = _require_text(
                    values[layout.raw_cell_id_column],
                    column=layout.raw_cell_id_column,
                    row_number=row_number,
                )
                nominal_capacity = _require_finite_float(
                    values[layout.nominal_capacity_ah_column],
                    column=layout.nominal_capacity_ah_column,
                    row_number=row_number,
                )
                if nominal_capacity <= 0:
                    raise ValueError(f"row {row_number} nominal capacity must be positive")
                prior_capacity = capacity_by_raw_id.setdefault(raw_cell_id, nominal_capacity)
                if not math.isclose(prior_capacity, nominal_capacity, rel_tol=0.0, abs_tol=1e-12):
                    raise ValueError(f"raw cell {raw_cell_id!r} has conflicting nominal capacities")

                protocol_id = (
                    _optional_text(
                        values[layout.protocol_id_column],
                        column=layout.protocol_id_column,
                        row_number=row_number,
                    )
                    if layout.protocol_id_column
                    else None
                )
                protocol_description = (
                    _optional_text(
                        values[layout.protocol_description_column],
                        column=layout.protocol_description_column,
                        row_number=row_number,
                    )
                    if layout.protocol_description_column
                    else None
                )
                prior_protocol = protocol_by_raw_id.setdefault(
                    raw_cell_id, (protocol_id, protocol_description)
                )
                if prior_protocol != (protocol_id, protocol_description):
                    raise ValueError(f"raw cell {raw_cell_id!r} has conflicting protocol metadata")

                records_by_raw_id.setdefault(raw_cell_id, []).append(
                    CycleRecord(
                        dataset_id=dataset_id,
                        cell_id=f"{dataset_id}_{raw_cell_id}",
                        cycle_index=_require_index(
                            values[layout.cycle_index_column],
                            column=layout.cycle_index_column,
                            row_number=row_number,
                        ),
                        sample_index=_require_index(
                            values[layout.sample_index_column],
                            column=layout.sample_index_column,
                            row_number=row_number,
                        ),
                        time_s=_require_finite_float(
                            values[layout.time_s_column],
                            column=layout.time_s_column,
                            row_number=row_number,
                        ),
                        voltage_v=_require_finite_float(
                            values[layout.voltage_v_column],
                            column=layout.voltage_v_column,
                            row_number=row_number,
                        ),
                        current_a=_require_finite_float(
                            values[layout.current_a_column],
                            column=layout.current_a_column,
                            row_number=row_number,
                        ),
                        temperature_c=(
                            _optional_float(
                                values[layout.temperature_c_column],
                                column=layout.temperature_c_column,
                                row_number=row_number,
                            )
                            if layout.temperature_c_column
                            else None
                        ),
                        charge_capacity_ah=(
                            _optional_float(
                                values[layout.charge_capacity_ah_column],
                                column=layout.charge_capacity_ah_column,
                                row_number=row_number,
                            )
                            if layout.charge_capacity_ah_column
                            else None
                        ),
                        discharge_capacity_ah=(
                            _optional_float(
                                values[layout.discharge_capacity_ah_column],
                                column=layout.discharge_capacity_ah_column,
                                row_number=row_number,
                            )
                            if layout.discharge_capacity_ah_column
                            else None
                        ),
                        internal_resistance_ohm=(
                            _optional_float(
                                values[layout.internal_resistance_ohm_column],
                                column=layout.internal_resistance_ohm_column,
                                row_number=row_number,
                            )
                            if layout.internal_resistance_ohm_column
                            else None
                        ),
                    )
                )
        finally:
            workbook.close()

    cells: list[NaumannCell] = []
    for raw_cell_id, records in records_by_raw_id.items():
        ordered_records = tuple(
            sorted(records, key=lambda record: (record.cycle_index, record.sample_index))
        )
        quality = validate_cycle_records(ordered_records)
        rejection_issues = tuple(
            issue
            for issue in quality.issues
            if issue.severity in {DataQualitySeverity.ERROR, DataQualitySeverity.BLOCKING}
        )
        if rejection_issues:
            codes = ", ".join(
                issue.code for issue in rejection_issues
            )
            raise ValueError(f"Naumann cell {raw_cell_id!r} failed structural validation: {codes}")
        samples_per_cycle: dict[int, int] = {}
        for record in ordered_records:
            samples_per_cycle[record.cycle_index] = (
                samples_per_cycle.get(record.cycle_index, 0) + 1
            )
        if not any(sample_count >= 2 for sample_count in samples_per_cycle.values()):
            raise ValueError(
                f"Naumann cell {raw_cell_id!r} does not contain a time-series cycle"
            )
        protocol_id, protocol_description = protocol_by_raw_id[raw_cell_id]
        cells.append(
            (
                CellMetadata(
                    dataset_id=dataset_id,
                    cell_id=f"{dataset_id}_{raw_cell_id}",
                    raw_cell_id=raw_cell_id,
                    chemistry="LFP/graphite",
                    nominal_capacity_ah=capacity_by_raw_id[raw_cell_id],
                    reference_capacity_ah=None,
                    protocol_id=protocol_id,
                    protocol_description=protocol_description,
                    official_life_label=None,
                    official_life_label_name=None,
                    source_uri=manifest.source_uri,
                    source_sha256=source_sha256,
                    schema_version="1.0",
                    adapter_version=_ADAPTER_VERSION,
                    ingestion_parameters={
                        "layout_version": layout.layout_version,
                        "sheet_name": layout.sheet_name,
                        "header_row": layout.header_row,
                        "units": layout.units.model_dump(exclude_none=True),
                        "columns": layout.model_dump(
                            exclude={"layout_version", "sheet_name", "header_row", "units"}
                        ),
                    },
                ),
                ordered_records,
            )
        )
    if not cells:
        raise ValueError("Naumann workbook contains no sample rows")
    return tuple(cells)
