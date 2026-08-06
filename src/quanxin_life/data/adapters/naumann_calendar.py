"""Explicit adapter for Naumann calendar-ageing workbook measurements.

The published workbooks contain condition-level storage-ageing measurements,
not a per-cell electrochemical time series.  They must therefore never be
coerced into :class:`CycleRecord` or used as a cell-level EOL80 training set.
This adapter emits only provenance-bound condition observations for the
calendar-ageing and operating-condition modelling path.
"""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Literal

from pydantic import Field, model_validator

from quanxin_life.core.schemas import ContractModel
from quanxin_life.data.manifest import RawFileManifest, Sha256, verify_raw_file
from quanxin_life.data.source_catalog import IngestionMode, SourceCatalogEntry

_ADAPTER_VERSION = "naumann-calendar-condition-v1.0.0"


class CalendarConditionColumn(ContractModel):
    """One reviewed condition column in a published calendar-ageing workbook."""

    column: int = Field(ge=1)
    expected_header: str = Field(min_length=1)
    condition_id: str = Field(min_length=1)
    temperature_c: float = Field(ge=-100, le=200, allow_inf_nan=False)
    mean_soc: float = Field(ge=0, le=1, allow_inf_nan=False)


class NaumannCalendarLayout(ContractModel):
    """Explicit workbook mapping; column semantics are never inferred from text."""

    layout_version: str = Field(min_length=1)
    identity_level: Literal["condition"] = "condition"
    cell_level_split_supported: Literal[False] = False
    sheet_name: str = Field(min_length=1)
    capacity_header_row: int = Field(ge=1)
    capacity_header_column: int = Field(ge=1)
    capacity_header: str = Field(min_length=1)
    time_header_row: int = Field(ge=1)
    first_observation_row: int = Field(ge=1)
    time_column: int = Field(ge=1)
    time_header: str = Field(min_length=1)
    condition_columns: tuple[CalendarConditionColumn, ...] = Field(min_length=1)

    @model_validator(mode="after")
    def columns_are_unique_and_ordered(self) -> NaumannCalendarLayout:
        if self.first_observation_row <= self.time_header_row:
            raise ValueError("first_observation_row must follow time_header_row")
        if self.capacity_header_row >= self.first_observation_row:
            raise ValueError("capacity_header_row must precede first_observation_row")
        if (
            self.capacity_header_row == self.time_header_row
            and self.capacity_header_column == self.time_column
        ):
            raise ValueError("capacity header cannot share the time header cell")
        columns = [condition.column for condition in self.condition_columns]
        condition_ids = [condition.condition_id for condition in self.condition_columns]
        if self.time_column in columns:
            raise ValueError("time_column cannot also be a condition column")
        if len(columns) != len(set(columns)):
            raise ValueError("calendar condition columns must be unique")
        if len(condition_ids) != len(set(condition_ids)):
            raise ValueError("calendar condition identifiers must be unique")
        return self


class CalendarCapacityObservation(ContractModel):
    """One directly measured capacity point under a declared calendar condition."""

    dataset_id: Literal["NAUMANN_CALENDAR"] = "NAUMANN_CALENDAR"
    condition_id: str = Field(min_length=1)
    storage_time_h: float = Field(ge=0, allow_inf_nan=False)
    temperature_c: float = Field(ge=-100, le=200, allow_inf_nan=False)
    mean_soc: float = Field(ge=0, le=1, allow_inf_nan=False)
    dod: None = None
    charge_c_rate: None = None
    discharge_c_rate: None = None
    capacity_ah: float = Field(ge=0, allow_inf_nan=False)
    source_file: str = Field(min_length=1)
    source_sha256: Sha256
    layout_version: str = Field(min_length=1)
    adapter_version: str = _ADAPTER_VERSION


def load_naumann_calendar_layout(path: Path) -> NaumannCalendarLayout:
    """Load a version-controlled, explicitly reviewed calendar layout JSON."""

    path = Path(path)
    if path.suffix.lower() != ".json":
        raise ValueError("calendar layout must use a .json file")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"calendar layout does not exist: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError("calendar layout must contain valid JSON") from exc
    return NaumannCalendarLayout.model_validate(payload)


def _require_finite_number(value: object, *, field: str, row: int) -> float:
    if isinstance(value, bool):
        raise ValueError(f"row {row} has invalid {field}")
    try:
        numeric = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError) as exc:
        raise ValueError(f"row {row} has invalid {field}") from exc
    if not math.isfinite(numeric):
        raise ValueError(f"row {row} has invalid {field}")
    return numeric


def _validated_layout(layout: NaumannCalendarLayout) -> NaumannCalendarLayout:
    """Revalidate to block caller-created Pydantic bypass objects."""

    return NaumannCalendarLayout.model_validate(layout.model_dump(mode="json"))


def _validate_source(
    path: Path,
    manifest: RawFileManifest,
    source: SourceCatalogEntry,
) -> None:
    if manifest.dataset_id != "NAUMANN_CALENDAR" or source.dataset_id != manifest.dataset_id:
        raise ValueError("calendar manifest and source dataset_id must be NAUMANN_CALENDAR")
    if source.ingestion_mode is not IngestionMode.TABULAR:
        raise ValueError("Naumann calendar source must use tabular ingestion")
    allowed_suffixes = {suffix.lower() for suffix in source.expected_suffixes}
    if path.suffix.lower() != ".xlsx" or ".xlsx" not in allowed_suffixes:
        raise ValueError("Naumann calendar workbook must use an approved .xlsx suffix")
    if manifest.source_uri != source.source_uri:
        raise ValueError("Naumann calendar manifest source URI does not match the source catalog")
    if manifest.license_name != source.license_status:
        raise ValueError("Naumann calendar manifest license does not match the source catalog")


def load_naumann_calendar_capacity(
    path: Path,
    manifest: RawFileManifest,
    source: SourceCatalogEntry,
    *,
    layout: NaumannCalendarLayout,
) -> tuple[CalendarCapacityObservation, ...]:
    """Load capacity measurements only from a provenance-verified workbook.

    The layout explicitly names each condition column.  No condition, unit or
    missing value is inferred, interpolated or silently imputed.
    """

    path = Path(path)
    _validate_source(path, manifest, source)
    layout = _validated_layout(layout)
    source_sha256 = verify_raw_file(path, manifest)

    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - only without optional data extras
        message = "Naumann calendar loading requires the 'data' optional dependencies"
        raise RuntimeError(message) from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if layout.sheet_name not in workbook.sheetnames:
            raise ValueError(f"calendar workbook is missing sheet {layout.sheet_name!r}")
        worksheet = workbook[layout.sheet_name]
        if (
            worksheet.cell(layout.capacity_header_row, layout.capacity_header_column).value
            != layout.capacity_header
        ):
            raise ValueError("calendar capacity header does not match reviewed layout")
        if worksheet.cell(layout.time_header_row, layout.time_column).value != layout.time_header:
            raise ValueError("calendar workbook time header does not match reviewed layout")
        for condition in layout.condition_columns:
            actual_header = worksheet.cell(layout.time_header_row, condition.column).value
            if actual_header != condition.expected_header:
                raise ValueError(
                    "calendar workbook expected_header mismatch for "
                    f"condition {condition.condition_id!r}"
                )

        observations: list[CalendarCapacityObservation] = []
        previous_time: float | None = None
        row_number = layout.first_observation_row
        while True:
            raw_time = worksheet.cell(row_number, layout.time_column).value
            raw_capacities = [
                worksheet.cell(row_number, condition.column).value
                for condition in layout.condition_columns
            ]
            if raw_time is None and all(value is None for value in raw_capacities):
                break
            if raw_time is None or any(value is None for value in raw_capacities):
                raise ValueError(f"row {row_number} has incomplete calendar measurement values")
            storage_time_h = _require_finite_number(raw_time, field="storage time", row=row_number)
            if previous_time is not None and storage_time_h <= previous_time:
                message = "calendar storage time must strictly increase without reordering rows"
                raise ValueError(message)
            previous_time = storage_time_h
            for condition, raw_capacity in zip(
                layout.condition_columns, raw_capacities, strict=True
            ):
                capacity_ah = _require_finite_number(
                    raw_capacity,
                    field=f"capacity for condition {condition.condition_id}",
                    row=row_number,
                )
                if capacity_ah < 0:
                    raise ValueError(f"row {row_number} has negative capacity")
                observations.append(
                    CalendarCapacityObservation(
                        condition_id=condition.condition_id,
                        storage_time_h=storage_time_h,
                        temperature_c=condition.temperature_c,
                        mean_soc=condition.mean_soc,
                        capacity_ah=capacity_ah,
                        source_file=manifest.relative_path,
                        source_sha256=source_sha256,
                        layout_version=layout.layout_version,
                    )
                )
            row_number += 1
    finally:
        workbook.close()

    if not observations:
        raise ValueError("calendar workbook contains no reviewed observations")
    return tuple(observations)
