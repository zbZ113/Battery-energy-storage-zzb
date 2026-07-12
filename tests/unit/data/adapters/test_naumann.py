import hashlib
from pathlib import Path
from typing import BinaryIO

import pytest
from pydantic import ValidationError

import quanxin_life.data.adapters.naumann as naumann_adapter
from quanxin_life.data.adapters.naumann import (
    NaumannWorkbookLayout,
    load_naumann_workbook,
)
from quanxin_life.data.manifest import RawFileManifest, verify_raw_file_stream
from quanxin_life.data.source_catalog import IngestionMode, SourceCatalogEntry


def _source(dataset_id: str = "NAUMANN_CYCLE", **updates: object) -> SourceCatalogEntry:
    values: dict[str, object] = {
        "dataset_id": dataset_id,
        "version": "Mendeley-v1",
        "source_uri": "https://data.mendeley.com/datasets/6hgyr25h8d/1",
        "paper_uri": "https://doi.org/10.1016/j.jpowsour.2019.227666",
        "license_status": "CC BY 4.0",
        "ingestion_mode": IngestionMode.TABULAR,
        "expected_suffixes": (".xlsx",),
    }
    values.update(updates)
    return SourceCatalogEntry.model_validate(values)


def _layout() -> NaumannWorkbookLayout:
    return NaumannWorkbookLayout(
        layout_version="naumann-controlled-fixture-v1",
        sheet_name="samples",
        header_row=1,
        raw_cell_id_column="cell",
        nominal_capacity_ah_column="nominal_ah",
        cycle_index_column="cycle",
        sample_index_column="sample",
        time_s_column="time_s",
        voltage_v_column="voltage_v",
        current_a_column="current_a",
        temperature_c_column="temperature_c",
        protocol_id_column="protocol_id",
        protocol_description_column="protocol_description",
        units={
            "nominal_capacity_ah": "Ah",
            "time_s": "s",
            "voltage_v": "V",
            "current_a": "A",
            "temperature_c": "degC",
        },
    )


def _write_workbook(path: Path, rows: list[list[object]]) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "samples"
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def _manifest(path: Path, **updates: object) -> RawFileManifest:
    values: dict[str, object] = {
        "dataset_id": "NAUMANN_CYCLE",
        "relative_path": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "source_uri": "https://data.mendeley.com/datasets/6hgyr25h8d/1",
        "license_name": "CC BY 4.0",
        "paper_doi": "10.1016/j.jpowsour.2019.227666",
    }
    values.update(updates)
    return RawFileManifest.model_validate(values)


def _rows() -> list[list[object]]:
    return [
        [
            "cell",
            "nominal_ah",
            "cycle",
            "sample",
            "time_s",
            "voltage_v",
            "current_a",
            "temperature_c",
            "protocol_id",
            "protocol_description",
        ],
        ["A01", 160.0, 1, 0, 0.0, 3.2, -80.0, 25.0, "T25-SOC50", "controlled fixture"],
        ["A01", 160.0, 1, 1, 10.0, 3.1, -80.0, 25.0, "T25-SOC50", "controlled fixture"],
        ["B02", 160.0, 1, 0, 0.0, 3.3, -80.0, None, "T25-SOC50", "controlled fixture"],
        ["B02", 160.0, 1, 1, 10.0, 3.2, -80.0, None, "T25-SOC50", "controlled fixture"],
    ]


def test_loads_explicit_layout_and_preserves_provenance(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())

    cells = load_naumann_workbook(path, _manifest(path), _source(), layout=_layout())

    assert [metadata.cell_id for metadata, _ in cells] == [
        "NAUMANN_CYCLE_A01",
        "NAUMANN_CYCLE_B02",
    ]
    metadata, records = cells[0]
    assert metadata.source_sha256 == _manifest(path).sha256
    assert metadata.reference_capacity_ah is None
    assert metadata.official_life_label is None
    assert metadata.ingestion_parameters["layout_version"] == "naumann-controlled-fixture-v1"
    assert [record.time_s for record in records] == [0.0, 10.0]
    assert cells[1][1][0].temperature_c is None


def test_rejects_hash_failure_before_opening_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())

    def workbook_open_is_forbidden(*args: object, **kwargs: object) -> object:
        raise AssertionError("workbook must not open before provenance verification")

    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", workbook_open_is_forbidden)
    with pytest.raises(ValueError, match="SHA-256 mismatch"):
        load_naumann_workbook(path, _manifest(path, sha256="0" * 64), _source(), layout=_layout())


def test_loads_workbook_from_the_same_stream_used_for_verification(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())
    manifest = _manifest(path)
    verifier_handle: object | None = None
    workbook_handle: object | None = None

    def verify_from_open_handle(
        handle: BinaryIO, workbook_path: Path, raw_manifest: RawFileManifest
    ) -> str:
        nonlocal verifier_handle
        assert workbook_path == path
        assert raw_manifest == manifest
        verifier_handle = handle
        return verify_raw_file_stream(handle, workbook_path, raw_manifest)

    def path_verification_is_forbidden(*args: object, **kwargs: object) -> str:
        raise AssertionError("Naumann loader must verify its already-open workbook stream")

    import openpyxl

    original_load_workbook = openpyxl.load_workbook

    def record_workbook_handle(
        handle: object, *args: object, **kwargs: object
    ) -> object:
        nonlocal workbook_handle
        workbook_handle = handle
        return original_load_workbook(handle, *args, **kwargs)

    monkeypatch.setattr(
        naumann_adapter, "verify_raw_file_stream", verify_from_open_handle, raising=False
    )
    monkeypatch.setattr(
        naumann_adapter, "verify_raw_file", path_verification_is_forbidden, raising=False
    )
    monkeypatch.setattr(openpyxl, "load_workbook", record_workbook_handle)

    cells = load_naumann_workbook(path, manifest, _source(), layout=_layout())

    assert cells
    assert verifier_handle is not None
    assert workbook_handle is verifier_handle


@pytest.mark.parametrize(
    "manifest_updates, source, match",
    [
        ({"dataset_id": "MATR"}, _source(), "dataset_id"),
        ({}, _source("NAUMANN_CALENDAR"), "dataset_id"),
        ({}, _source(source_uri="https://example.invalid"), "source URI"),
    ],
)
def test_rejects_dataset_and_catalog_mismatches(
    tmp_path: Path,
    manifest_updates: dict[str, object],
    source: SourceCatalogEntry,
    match: str,
) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())

    with pytest.raises(ValueError, match=match):
        load_naumann_workbook(path, _manifest(path, **manifest_updates), source, layout=_layout())


def test_rejects_missing_sheet_or_required_column(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())

    with pytest.raises(ValueError, match="sheet"):
        load_naumann_workbook(
            path,
            _manifest(path),
            _source(),
            layout=_layout().model_copy(update={"sheet_name": "unknown"}),
        )

    incomplete = [
        [value for index, value in enumerate(row) if index != 6]
        for row in _rows()
    ]
    _write_workbook(path, incomplete)
    with pytest.raises(ValueError, match="required columns"):
        load_naumann_workbook(path, _manifest(path), _source(), layout=_layout())


def test_rejects_duplicate_sample_keys_and_noncanonical_units(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    rows = [
        *_rows(),
        ["A01", 160.0, 1, 1, 20.0, 3.0, -80.0, 25.0, "T25-SOC50", "controlled fixture"],
    ]
    _write_workbook(path, rows)

    with pytest.raises(ValueError, match="DUPLICATE_SAMPLE"):
        load_naumann_workbook(path, _manifest(path), _source(), layout=_layout())

    with pytest.raises(ValueError, match="voltage_v"):
        NaumannWorkbookLayout.model_validate(
            {
                **_layout().model_dump(),
                "units": {**_layout().units.model_dump(), "voltage_v": "mV"},
            }
        )


def test_units_are_deeply_immutable() -> None:
    layout = _layout()

    with pytest.raises(ValidationError, match="frozen"):
        layout.units.voltage_v = "mV"  # type: ignore[attr-defined]


def test_loader_revalidates_units_modified_through_model_copy(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())
    bypassed_layout = _layout().model_copy(
        update={
            "units": {
                "nominal_capacity_ah": "Ah",
                "time_s": "s",
                "voltage_v": "mV",
                "current_a": "A",
                "temperature_c": "degC",
            }
        }
    )

    with pytest.raises(ValueError, match="voltage_v"):
        load_naumann_workbook(path, _manifest(path), _source(), layout=bypassed_layout)


def test_loader_revalidates_duplicate_columns_modified_through_model_copy(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    _write_workbook(path, _rows())
    bypassed_layout = _layout().model_copy(update={"time_s_column": "cycle"})

    with pytest.raises(ValueError, match="layout column names must be unique"):
        load_naumann_workbook(path, _manifest(path), _source(), layout=bypassed_layout)


def test_rejects_non_monotonic_time_series(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    rows = _rows()
    rows[2][4] = 0.0
    _write_workbook(path, rows)

    with pytest.raises(ValueError, match="NON_MONOTONIC_TIME"):
        load_naumann_workbook(path, _manifest(path), _source(), layout=_layout())


def test_rejects_summary_only_data_without_a_time_series_cycle(tmp_path: Path) -> None:
    path = tmp_path / "naumann_cycle.xlsx"
    rows = [
        _rows()[0],
        ["A01", 160.0, 1, 0, 0.0, 3.2, -80.0, 25.0, "T25-SOC50", "controlled fixture"],
        ["A01", 160.0, 2, 0, 0.0, 3.1, -80.0, 25.0, "T25-SOC50", "controlled fixture"],
    ]
    _write_workbook(path, rows)

    with pytest.raises(ValueError, match="time-series cycle"):
        load_naumann_workbook(path, _manifest(path), _source(), layout=_layout())
